# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Alexandre Tournel

"""Tests de la détection de titres par mise en forme (police+gras) et du nettoyage
du filigrane/pied de page dans l'extraction PDF et DOCX (`pipeline/extract.py`).
Fixtures construites directement avec `fitz`/`python-docx` (déjà des dépendances) —
pas besoin de fichiers de production réels pour ces tests unitaires."""
import re
from pathlib import Path

import docx
import fitz
from docx.shared import Pt

from pipeline.extract import (
    _docx_heading_candidates,
    _extract_docx,
    _extract_pdf,
    _normalize_for_match,
    _promote_docx_headings,
)

BODY_SIZE = 11
PART_SIZE = 16   # palier "Partie" (le plus fréquent, le plus petit des 2 titres)
CHAP_SIZE = 22   # palier "Chapitre" (le plus large, le moins fréquent)


def _add_page(doc, lines):
    """lines: liste de (y, texte, taille, gras). Une page 400x600pt."""
    page = doc.new_page(width=400, height=600)
    for y, text, size, bold in lines:
        page.insert_text((40, y), text, fontsize=size, fontname="hebo" if bold else "helv")
    return page


def _make_pdf(pages: list[list[tuple]], tmp_path) -> Path:
    doc = fitz.open()
    for lines in pages:
        _add_page(doc, lines)
    path = tmp_path / "fixture.pdf"
    doc.save(str(path))
    doc.close()
    return path


def _headings(text: str) -> list[tuple[str, str]]:
    return re.findall(r"(?m)^(#{1,6})\s+(.*)$", text)


def test_two_tiers_promoted_to_hash_levels(tmp_path):
    pages = [
        [(60, "Chapter One", CHAP_SIZE, True),
         (110, "Part A", PART_SIZE, True),
         (150, "Some ordinary body text describing the scene in plain prose.", BODY_SIZE, False)],
        [(60, "Part B", PART_SIZE, True),
         (110, "More ordinary narration text continues here for the reader.", BODY_SIZE, False)],
    ]
    path = _make_pdf(pages, tmp_path)
    res = _extract_pdf(path, tmp_path / "media", extract_images=False)
    heads = _headings(res.text)
    assert ("#", "Chapter One") in heads
    assert ("##", "Part A") in heads
    assert ("##", "Part B") in heads


def test_body_text_never_promoted(tmp_path):
    pages = [[(60, "Chapter One", CHAP_SIZE, True),
              (110, "Just plain narration, never a heading candidate here.", BODY_SIZE, False)]]
    path = _make_pdf(pages, tmp_path)
    res = _extract_pdf(path, tmp_path / "media", extract_images=False)
    assert "Just plain narration" in res.text
    assert "# Just plain narration" not in res.text


def test_repeated_footer_stripped_from_body(tmp_path):
    pages = []
    for i in range(5):
        pages.append([
            (60, "Chapter One", CHAP_SIZE, True),
            (110, f"Narration paragraph number {i} continues the story onward.", BODY_SIZE, False),
            (580, f"Page {i + 1} | scanlation-credit.example", 9, True),  # bande basse de page
        ])
    path = _make_pdf(pages, tmp_path)
    res = _extract_pdf(path, tmp_path / "media", extract_images=False)
    assert "scanlation-credit.example" not in res.text
    assert "Page 1" not in res.text
    assert "Narration paragraph number 0" in res.text


def test_denylisted_title_not_promoted_but_kept_as_text(tmp_path):
    pages = [[(60, "Table of Contents", CHAP_SIZE, True),
              (110, "1. Chapter One\n2. Chapter Two", BODY_SIZE, False)]]
    path = _make_pdf(pages, tmp_path)
    res = _extract_pdf(path, tmp_path / "media", extract_images=False)
    assert "Table of Contents" in res.text
    assert "# Table of Contents" not in res.text


def test_repeated_decorative_subtitle_at_tier_not_promoted(tmp_path):
    """Un sous-titre décoratif identique répété au même palier (ex. le nom de la
    série redonné à chaque chapitre) ne doit PAS devenir un titre de Partie —
    contrairement à un vrai titre de Partie unique."""
    pages = [
        [(60, "Chapter One", CHAP_SIZE, True),
         (110, "My Book Subtitle", PART_SIZE, True),
         (150, "Part Alpha", PART_SIZE, True),
         (190, "Narration text for the first chapter goes here today.", BODY_SIZE, False)],
        [(60, "Chapter Two", CHAP_SIZE, True),
         (110, "My Book Subtitle", PART_SIZE, True),
         (150, "Part Beta", PART_SIZE, True),
         (190, "Narration text for the second chapter goes here today.", BODY_SIZE, False)],
    ]
    path = _make_pdf(pages, tmp_path)
    res = _extract_pdf(path, tmp_path / "media", extract_images=False)
    heads = _headings(res.text)
    titles = [t for _, t in heads]
    assert "Part Alpha" in titles and "Part Beta" in titles
    assert "My Book Subtitle" not in titles
    assert "My Book Subtitle" in res.text  # reste en texte simple, pas supprimé


# --------------------------------------------------------------------------- #
#  DOCX
# --------------------------------------------------------------------------- #
# Taille explicite sur chaque run de corps de texte : reflète les vrais .docx du
# projet (vérifié sur roman B Vol.1 FR — Word écrit systématiquement une taille
# explicite par run), plutôt que le modèle vierge de python-docx dont le style
# « Normal » n'expose aucune taille via l'API (docDefaults bruts, hors périmètre).
DOCX_BODY_PT = 11
DOCX_PART_PT = 16
DOCX_CHAP_PT = 22


def _body_paragraph(doc, text="Ordinary narration text used to establish body size here.", pt=DOCX_BODY_PT):
    p = doc.add_paragraph()
    p.add_run(text).font.size = Pt(pt)
    return p


def _bold_paragraph(doc, text, pt):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(pt)
    return p


def _make_docx(tmp_path, builder) -> Path:
    d = docx.Document()
    builder(d)
    path = tmp_path / "fixture.docx"
    d.save(str(path))
    return path


def test_docx_heading_candidates_two_bold_tiers_levels(tmp_path):
    """Aucun style Word nommé : la hiérarchie vient PUREMENT de la mise en forme
    (2 tailles de gras distinctes au-dessus du corps de texte)."""
    def build(d):
        _bold_paragraph(d, "Chapter One", DOCX_CHAP_PT)
        _bold_paragraph(d, "Part Alpha", DOCX_PART_PT)
        _body_paragraph(d)
        _body_paragraph(d)
        _bold_paragraph(d, "Part Beta", DOCX_PART_PT)
        _body_paragraph(d)

    path = _make_docx(tmp_path, build)
    candidates = _docx_heading_candidates(path)
    assert candidates == [("Chapter One", 1), ("Part Alpha", 2), ("Part Beta", 2)]


def test_docx_heading_candidates_named_style_included_in_order(tmp_path):
    """Un titre déjà stylé (Heading 1) doit apparaître dans les candidats, à sa
    place dans l'ordre du document — nécessaire pour que le curseur de
    `_promote_docx_headings` avance correctement au-delà, même si ce signal est
    redondant avec ce que Pandoc produit déjà pour ce paragraphe précis."""
    def build(d):
        d.add_heading("Chapter One", level=1)
        _body_paragraph(d)

    path = _make_docx(tmp_path, build)
    candidates = _docx_heading_candidates(path)
    assert candidates == [("Chapter One", 1)]


def test_docx_heading_candidates_excludes_denylisted_title(tmp_path):
    def build(d):
        _bold_paragraph(d, "Table of Contents", 20)
        _body_paragraph(d)

    path = _make_docx(tmp_path, build)
    candidates = _docx_heading_candidates(path)
    assert candidates == []


def test_docx_heading_candidates_ignores_plain_body_text(tmp_path):
    def build(d):
        for _ in range(3):
            _body_paragraph(d)

    path = _make_docx(tmp_path, build)
    assert _docx_heading_candidates(path) == []


def test_promote_docx_headings_sequential_match_and_skip_unmatched():
    md = "Some intro line.\n\nChapter One\n\nBody text continues here.\n\nPart Alpha\n\nMore body."
    candidates = [("Chapter One", 1), ("Not Found Anywhere", 2), ("Part Alpha", 2)]
    out = _promote_docx_headings(md, candidates)
    lines = out.splitlines()
    assert "# Chapter One" in lines
    assert "## Part Alpha" in lines
    assert "Not Found Anywhere" not in out  # jamais planté, simplement ignoré


def test_promote_docx_headings_idempotent_on_existing_atx_line():
    md = "## Already A Heading\n\nBody text."
    out = _promote_docx_headings(md, [("Already A Heading", 1)])
    assert out.count("#") == 2  # pas de double préfixe ("# ## ...")


def test_normalize_for_match_neutralizes_pandoc_escaping_and_quotes():
    assert _normalize_for_match("## Chapter\\*One") == _normalize_for_match("chapter*one")
    assert _normalize_for_match("Chapter’s End") == _normalize_for_match("chapter's end")


def test_extract_docx_end_to_end_promotes_bold_unstyled_heading(tmp_path):
    """Bout en bout via Pandoc réel : des titres tapés en gras/grand SANS style Word
    nommé doivent ressortir comme de vrais titres ATX dans le markdown final."""
    def build(d):
        _bold_paragraph(d, "Chapter One", DOCX_CHAP_PT)
        _bold_paragraph(d, "Part Alpha", DOCX_PART_PT)
        _body_paragraph(d, "Some narration paragraph goes here for the reader to enjoy.")
        _bold_paragraph(d, "Part Beta", DOCX_PART_PT)
        _body_paragraph(d, "Another narration paragraph follows right after that one.")

    path = _make_docx(tmp_path, build)
    res = _extract_docx(path, tmp_path / "media", extract_images=False)
    heads = _headings(res.text)
    assert ("#", "Chapter One") in heads
    assert ("##", "Part Alpha") in heads
    assert ("##", "Part Beta") in heads


def test_extract_docx_end_to_end_keeps_named_heading_style(tmp_path):
    """Un titre déjà stylé (Heading 1) continue de fonctionner comme avant (Pandoc
    seul suffit) — la nouvelle passe ne doit rien casser sur ce chemin existant."""
    def build(d):
        d.add_heading("Chapter One", level=1)
        _body_paragraph(d, "Narration text following the styled heading paragraph.")

    path = _make_docx(tmp_path, build)
    res = _extract_docx(path, tmp_path / "media", extract_images=False)
    assert ("#", "Chapter One") in _headings(res.text)
